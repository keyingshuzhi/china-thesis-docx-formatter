"""Local MCP server for the academic paper formatter."""
from __future__ import annotations

import base64
import json
import logging
import os
import sys
import tempfile
import zipfile
from collections.abc import Callable
from contextlib import contextmanager
from pathlib import Path
from typing import Annotated, Any, Literal
from urllib.parse import urlsplit
from urllib.request import url2pathname

from docx import Document
from lxml import etree
from mcp.server.mcpserver import Context, ListRoots, MCPServer, Resolve
from mcp.types import (
    BlobResourceContents,
    CallToolResult,
    EmbeddedResource,
    ListRootsResult,
    TextContent,
    TextResourceContents,
    ToolAnnotations,
)

from . import __version__

DEFAULT_REPOSITORY_ROOT = Path(__file__).resolve().parents[3]
REPOSITORY_ROOT = Path(
    os.environ.get("THESIS_FORMATTER_REPO_ROOT", DEFAULT_REPOSITORY_ROOT)
).expanduser().resolve()
SERVER_NAME = "academic-paper-formatter"
SCRIPTS_DIR = REPOSITORY_ROOT / "scripts"
if not SCRIPTS_DIR.is_dir():
    raise RuntimeError(
        "The MCP server must run from the project checkout; "
        "set THESIS_FORMATTER_REPO_ROOT to the repository root."
    )
sys.path.insert(0, str(SCRIPTS_DIR))

# These are the same operations used by the existing CLI entry points.
from complete_content import find_placeholders, paragraphs, replace_in_paragraph  # noqa: E402
from docx_rules import extract_rules, read_rules  # noqa: E402
from format_document import format_document as apply_document_format  # noqa: E402
from manage_references import (  # noqa: E402
    audit as audit_reference_document,
    insert_authoritative_entries,
    renumber as renumber_reference_labels,
)
from normalize_floating_figures import NS as FIGURE_NS  # noqa: E402
from normalize_floating_figures import configure_anchor  # noqa: E402
from parse_pdf_rules import extract_pdf_rules  # noqa: E402
from update_fields import field_counts  # noqa: E402
from validate_format import (  # noqa: E402
    audit_degree_structure,
    audit_document,
    validate_rules as check_rules,
)

logger = logging.getLogger("academic_paper_formatter_mcp")
Degree = Literal["auto", "undergraduate", "master", "doctoral"]
Horizontal = Literal["left", "center", "right"]
Vertical = Literal["top", "center", "bottom"]
RelativeTo = Literal["paragraph", "margin"]

server = MCPServer(
    name=SERVER_NAME,
    title="学位论文格式化",
    description="依据学校模板分析、格式化并审计中国学位论文 DOCX 文件。",
    version=__version__,
    instructions=(
        "始终以学校 DOCX 模板和书面规范为格式依据。保留原始文件，将结果写入单独路径，"
        "并将规则提取警告交由人工复核。自动使用当前 MCP 客户端提供的 Agent 工作区；"
        "若客户端没有提供 Roots，必须从当前 Agent 会话上下文取得工作区绝对路径，并在每次文件工具调用中传入 workspace_root；"
        "该参数不会跨调用缓存。THESIS_FORMATTER_WORKSPACE_ROOT 仅用于旧客户端固定目录兼容。"
        "依赖前一步结果的文件工具必须串行调用，并等待前一步返回 ok=true 后再继续；函数调用请求文本不代表工具已执行。"
        "若写入工具返回输出已存在，不得用同一路径重试或自动启用覆盖；先只读验证现有文件，或经确认后使用新输出路径。"
        "需要用户在会话中直接打开的最终产物，在完成校验后调用 load_generated_artifact，将工作区内文件作为嵌入资源返回；不要加载输入文件或内部中间文件。"
    ),
)
_READ_ONLY = ToolAnnotations(read_only_hint=True)
_WRITE_OUTPUT = ToolAnnotations(
    read_only_hint=False,
    destructive_hint=True,
    idempotent_hint=False,
    open_world_hint=False,
)
_ARTIFACT_MIME_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
    ".json": "application/json",
    ".md": "text/markdown",
    ".txt": "text/plain",
}
_TEXT_ARTIFACT_SUFFIXES = {".json", ".md", ".txt"}
_MAX_ARTIFACT_BYTES = 32 * 1024 * 1024


def _is_within(path: Path, root: Path) -> bool:
    return path == root or root in path.parents


def _validate_workspace_root(candidate: Path) -> Path:
    if not candidate.is_absolute():
        raise ValueError("The Agent workspace path must be absolute.")
    root = candidate.resolve(strict=True)
    if not root.is_dir():
        raise ValueError(f"Agent workspace root is not a directory: {root}")
    if _is_within(root, REPOSITORY_ROOT) or _is_within(REPOSITORY_ROOT, root):
        raise ValueError(
            "The Agent workspace must be separate from the source checkout and cannot be its parent directory."
        )
    return root


async def _request_workspace_source(ctx: Context) -> ListRoots | list[str]:
    capabilities = ctx.client_capabilities
    if capabilities is not None and capabilities.roots is not None:
        return ListRoots()
    return []


def _resolve_workspace_roots(
    source: Annotated[ListRootsResult | list[str], Resolve(_request_workspace_source)],
) -> list[Path]:
    candidates: list[Path] = []
    if isinstance(source, ListRootsResult):
        for root in source.roots:
            parsed = urlsplit(str(root.uri))
            if parsed.scheme != "file" or parsed.netloc not in ("", "localhost"):
                continue
            candidates.append(Path(url2pathname(parsed.path)))
    else:
        candidates.extend(Path(item).expanduser() for item in source if item.strip())

    roots: list[Path] = []
    for candidate in candidates:
        try:
            root = _validate_workspace_root(candidate)
        except (OSError, ValueError):
            continue
        if root not in roots:
            roots.append(root)

    return roots


def _effective_workspace_roots(
    client_roots: list[Path], workspace_root: str | None
) -> tuple[list[Path], str]:
    """Resolve the effective boundary without persisting one workspace across calls."""

    if workspace_root is not None:
        if not workspace_root.strip():
            raise ValueError("workspace_root must be a non-empty absolute Agent workspace path.")
        hinted_root = _validate_workspace_root(Path(workspace_root).expanduser())
        if client_roots:
            if not any(_is_within(hinted_root, root) for root in client_roots):
                raise ValueError(
                    "workspace_root must be within one of the MCP client-provided workspace Roots."
                )
        return [hinted_root], "workspace_hint"

    if client_roots:
        return client_roots, "client_roots"

    configured = os.environ.get("THESIS_FORMATTER_WORKSPACE_ROOT")
    if configured and configured.strip():
        return [_validate_workspace_root(Path(configured).expanduser())], "environment"

    return [], "unavailable"


def _require_workspace_roots(workspace_roots: list[Path]) -> list[Path]:
    if not workspace_roots:
        raise ValueError(
            "No Agent workspace is available. Pass the current Agent workspace as workspace_root, "
            "enable MCP Roots support, or configure THESIS_FORMATTER_WORKSPACE_ROOT as a fixed fallback."
        )
    return workspace_roots


def _require_allowed(path: Path, workspace_roots: list[Path]) -> Path:
    if not any(_is_within(path, root) for root in workspace_roots):
        raise ValueError(f"Path is outside the Agent workspace: {path}")
    return path


def _input_path(value: str, workspace_roots: list[Path], suffix: str | None = None) -> Path:
    workspace_roots = _require_workspace_roots(workspace_roots)
    candidate = Path(value).expanduser()
    if not candidate.is_absolute():
        if len(workspace_roots) != 1:
            raise ValueError("Relative paths require exactly one Agent workspace root; use an absolute path instead.")
        candidate = workspace_roots[0] / candidate

    # Check the resolved boundary before probing file metadata or suffix. This
    # makes every out-of-workspace input fail for the same reason, including a
    # path that is missing or has an unexpected extension, while the second
    # check below still blocks symlinks that resolve outside the workspace.
    _require_allowed(candidate.resolve(strict=False), workspace_roots)
    path = candidate.resolve(strict=True)
    _require_allowed(path, workspace_roots)
    if not path.is_file():
        raise ValueError(f"Expected a file: {path}")
    if suffix and path.suffix.casefold() != suffix:
        raise ValueError(f"Expected a {suffix} file: {path}")
    return path


def _output_path(
    value: str,
    *,
    workspace_roots: list[Path],
    inputs: tuple[Path, ...] = (),
    suffix: str | None = None,
    overwrite: bool = False,
) -> Path:
    workspace_roots = _require_workspace_roots(workspace_roots)
    candidate = Path(value).expanduser()
    if not candidate.is_absolute():
        if len(workspace_roots) != 1:
            raise ValueError("Relative paths require exactly one Agent workspace root; use an absolute path instead.")
        candidate = workspace_roots[0] / candidate
    if not candidate.name:
        raise ValueError("The output path must name a file.")
    parent_candidate = candidate.parent
    resolved_parent = parent_candidate.resolve(strict=False)
    _require_allowed(resolved_parent, workspace_roots)
    parent_candidate.mkdir(parents=True, exist_ok=True)
    parent = parent_candidate.resolve(strict=True)
    if not parent.is_dir():
        raise ValueError(f"Output parent is not a directory: {parent}")
    path = (parent / candidate.name).resolve(strict=False)
    if suffix and path.suffix.casefold() != suffix:
        raise ValueError(f"Expected an output file ending in {suffix}: {path}")
    _require_allowed(path, workspace_roots)
    if any(path == source or (path.exists() and source.exists() and path.samefile(source)) for source in inputs):
        raise ValueError("The output path must be different from every input file.")
    if path.exists() and not overwrite:
        raise FileExistsError(f"Output already exists; choose another path or explicitly set overwrite=true: {path}")
    if path.exists() and not path.is_file():
        raise ValueError(f"Output path is not a file: {path}")
    return path


@contextmanager
def _atomic_output(path: Path):
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    os.close(file_descriptor)
    temporary = Path(temporary_name)
    try:
        yield temporary
        if not temporary.is_file():
            raise RuntimeError("The formatter did not produce the expected output file.")
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    with _atomic_output(path) as temporary:
        temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _save_document(document: Any, path: Path) -> None:
    with _atomic_output(path) as temporary:
        document.save(temporary)


def _execute(operation: str, action: Callable[[], dict[str, Any]]) -> dict[str, Any]:
    try:
        return {"ok": True, **action()}
    except Exception as error:
        logger.exception("%s failed", operation)
        return {"ok": False, "error": str(error)}


@server.tool(title="MCP 服务版本", annotations=_READ_ONLY)
def get_mcp_version(
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """返回服务版本及工作区状态；workspace_root 可提供本次调用的动态工作区提示。"""

    try:
        effective_roots, workspace_source = _effective_workspace_roots(
            workspace_roots, workspace_root
        )
        workspace_error = None
    except (OSError, ValueError) as error:
        effective_roots, workspace_source = [], "invalid_workspace_hint"
        workspace_error = str(error)

    workspace: dict[str, Any] = {
        "configured": bool(effective_roots),
        "roots": [str(root) for root in effective_roots],
        "source": workspace_source,
    }
    if workspace_error:
        workspace["error"] = workspace_error
    elif not effective_roots:
        workspace["error"] = (
            "No Agent workspace was provided. Pass the current Agent workspace as workspace_root, "
            "or enable MCP Roots/workspace support."
        )
    return {
        "server_name": SERVER_NAME,
        "version": __version__,
        "python_version": ".".join(map(str, sys.version_info[:3])),
        "workspace": workspace,
    }


@server.tool(title="加载生成的文件", annotations=_READ_ONLY)
def load_generated_artifact(
    artifact_path: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    workspace_root: str | None = None,
) -> CallToolResult:
    """将 Agent 工作区中的文件作为 MCP 嵌入资源返回，供客户端在会话中呈现或打开。"""

    try:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        path = _input_path(artifact_path, effective_roots)
        suffix = path.suffix.casefold()
        mime_type = _ARTIFACT_MIME_TYPES.get(suffix)
        if mime_type is None:
            raise ValueError(
                "Unsupported artifact type. Supported files are .docx, .pdf, .json, .md, and .txt."
            )

        if path.stat().st_size > _MAX_ARTIFACT_BYTES:
            raise ValueError(
                "Artifact is larger than the 32 MiB session-loading limit."
            )
        with path.open("rb") as artifact_file:
            data = artifact_file.read(_MAX_ARTIFACT_BYTES + 1)
        if len(data) > _MAX_ARTIFACT_BYTES:
            raise ValueError("Artifact is larger than the 32 MiB session-loading limit.")
        size_bytes = len(data)

        uri = path.as_uri()
        if suffix in _TEXT_ARTIFACT_SUFFIXES:
            resource = TextResourceContents(
                uri=uri,
                mime_type=mime_type,
                text=data.decode("utf-8"),
            )
        else:
            resource = BlobResourceContents(
                uri=uri,
                mime_type=mime_type,
                blob=base64.b64encode(data).decode("ascii"),
            )

        metadata = {
            "ok": True,
            "artifact": {
                "name": path.name,
                "path": str(path),
                "uri": uri,
                "mime_type": mime_type,
                "size_bytes": size_bytes,
                "delivery": "embedded_resource",
            },
        }
        return CallToolResult(
            content=[
                TextContent(
                    text=f"已将文件内容作为 MCP 嵌入资源返回：{path.name}（{size_bytes} bytes）。"
                ),
                EmbeddedResource(resource=resource),
            ],
            structured_content=metadata,
        )
    except Exception as error:
        logger.exception("load_generated_artifact failed")
        return CallToolResult(
            content=[TextContent(text=f"加载生成文件失败：{error}")],
            structured_content={"ok": False, "error": str(error)},
            is_error=True,
        )


def _parse_xml(data: bytes):
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    return etree.fromstring(data, parser=parser)


def _rewrite_docx(
    source: Path,
    output: Path,
    transform: Callable[[str, bytes], bytes],
) -> None:
    with _atomic_output(output) as temporary:
        with zipfile.ZipFile(source, "r") as archive, zipfile.ZipFile(
            temporary, "w", compression=zipfile.ZIP_DEFLATED
        ) as result:
            for entry in archive.infolist():
                result.writestr(entry, transform(entry.filename, archive.read(entry)))


@server.tool(annotations=_WRITE_OUTPUT)
def analyze_template(
    template_path: str,
    output_path: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    degree: Degree = "auto",
    overwrite: bool = False,
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """从学校 DOCX 模板提取格式规则到新 JSON 文件，并返回待人工复核的提示。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        template = _input_path(template_path, effective_roots, ".docx")
        output = _output_path(
            output_path,
            workspace_roots=effective_roots,
            inputs=(template,),
            suffix=".json",
            overwrite=overwrite,
        )
        rules = extract_rules(Document(template), template.name, degree)
        _write_json(output, rules)
        return {
            "rules_path": str(output),
            "degree": degree,
            "style_count": len(rules.get("styles", {})),
            "extraction_warnings": rules.get("extraction_warnings", []),
            "document_structure": rules.get("document_structure", {}),
        }

    return _execute("analyze_template", action)


@server.tool(annotations=_READ_ONLY)
def validate_rules_file(
    rules_path: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """校验模板规则 JSON 的结构与必需字段，不修改文件。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        path = _input_path(rules_path, effective_roots, ".json")
        rules = read_rules(path)
        errors = check_rules(rules)
        return {"rules_path": str(path), "valid": not errors, "errors": errors}

    return _execute("validate_rules_file", action)


@server.tool(annotations=_WRITE_OUTPUT)
def format_document_copy(
    source_docx: str,
    rules_json: str,
    output_docx: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    overwrite: bool = False,
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """依据学校模板规则格式化 DOCX 副本，不修改源文件。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        source = _input_path(source_docx, effective_roots, ".docx")
        rules_path = _input_path(rules_json, effective_roots, ".json")
        output = _output_path(
            output_docx,
            workspace_roots=effective_roots,
            inputs=(source, rules_path),
            suffix=".docx",
            overwrite=overwrite,
        )
        rules = read_rules(rules_path)
        errors = check_rules(rules)
        if errors:
            raise ValueError("Rules are invalid: " + "; ".join(errors))
        document = Document(source)
        counts = apply_document_format(document, rules)
        _save_document(document, output)
        return {"output_path": str(output), "applied_counts": counts}

    return _execute("format_document_copy", action)


@server.tool(annotations=_READ_ONLY)
def validate_document_copy(
    rules_json: str,
    document_path: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    degree: Degree | None = None,
    strict_degree: bool = False,
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """按提取的规则审计 DOCX，可附带学位论文结构复核提示，不修改文件。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        rules_path = _input_path(rules_json, effective_roots, ".json")
        document_file = _input_path(document_path, effective_roots, ".docx")
        rules = read_rules(rules_path)
        rule_errors = check_rules(rules)
        if strict_degree and degree is None:
            raise ValueError("degree is required when strict_degree=true.")
        format_findings = []
        degree_warnings = []
        if not rule_errors:
            document = Document(document_file)
            format_findings = audit_document(document, rules)
            if degree:
                degree_warnings = audit_degree_structure(document, degree)
        valid = not rule_errors and not format_findings and not (strict_degree and degree_warnings)
        return {
            "document_path": str(document_file),
            "valid": valid,
            "rule_errors": rule_errors,
            "format_findings": format_findings,
            "degree_review_warnings": degree_warnings,
            "strict_degree": strict_degree,
        }

    return _execute("validate_document_copy", action)


@server.tool(annotations=_WRITE_OUTPUT)
def parse_pdf_guide(
    guide_pdf: str,
    output_json: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    overwrite: bool = False,
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """从学校 PDF 指南提取带页码的文本证据，供人工复核。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        source = _input_path(guide_pdf, effective_roots, ".pdf")
        output = _output_path(
            output_json,
            workspace_roots=effective_roots,
            inputs=(source,),
            suffix=".json",
            overwrite=overwrite,
        )
        payload = extract_pdf_rules(source)
        _write_json(output, payload)
        return {
            "output_path": str(output),
            "page_count": len(payload["pages"]),
            "requirement_count": len(payload["requirements"]),
            "extraction_warnings": payload["extraction_warnings"],
        }

    return _execute("parse_pdf_guide", action)


@server.tool(annotations=_READ_ONLY)
def inspect_placeholders(
    document_path: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """查找 DOCX 中的内容占位符，并返回位置和简短上下文，不修改文件。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        source = _input_path(document_path, effective_roots, ".docx")
        found = find_placeholders(Document(source))
        return {"document_path": str(source), "placeholder_count": len(found), "placeholders": found}

    return _execute("inspect_placeholders", action)


@server.tool(annotations=_WRITE_OUTPUT)
def apply_approved_replacements(
    source_docx: str,
    output_docx: str,
    replacements: dict[str, str],
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    overwrite: bool = False,
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """仅将用户审核的占位符替换映射应用到 DOCX 副本。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        source = _input_path(source_docx, effective_roots, ".docx")
        output = _output_path(
            output_docx,
            workspace_roots=effective_roots,
            inputs=(source,),
            suffix=".docx",
            overwrite=overwrite,
        )
        if not replacements:
            raise ValueError("Provide the exact, user-approved replacement mapping.")
        if not all(isinstance(key, str) and isinstance(value, str) for key, value in replacements.items()):
            raise ValueError("replacements must map strings to strings.")
        document = Document(source)
        applied = 0
        skipped: list[str] = []
        for paragraph in paragraphs(document):
            count, unresolved = replace_in_paragraph(paragraph, replacements)
            applied += count
            skipped.extend(unresolved)
        _save_document(document, output)
        return {
            "output_path": str(output),
            "applied_replacements": applied,
            "split_run_placeholders": sorted(set(skipped)),
        }

    return _execute("apply_approved_replacements", action)


@server.tool(annotations=_READ_ONLY)
def audit_references(
    source_docx: str,
    rules_json: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """审计正文数字引用与参考文献条目，不修改 DOCX。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        source = _input_path(source_docx, effective_roots, ".docx")
        rules_path = _input_path(rules_json, effective_roots, ".json")
        return {
            "document_path": str(source),
            "audit": audit_reference_document(Document(source), read_rules(rules_path)),
        }

    return _execute("audit_references", action)


@server.tool(annotations=_WRITE_OUTPUT)
def update_references_copy(
    source_docx: str,
    rules_json: str,
    output_docx: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    approved_metadata: dict[str, str] | None = None,
    renumber_existing: bool = False,
    overwrite: bool = False,
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """仅插入用户提供的权威书目数据，或在审计通过后重排参考文献编号，并写入副本。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        source = _input_path(source_docx, effective_roots, ".docx")
        rules_path = _input_path(rules_json, effective_roots, ".json")
        output = _output_path(
            output_docx,
            workspace_roots=effective_roots,
            inputs=(source, rules_path),
            suffix=".docx",
            overwrite=overwrite,
        )
        if approved_metadata is None and not renumber_existing:
            raise ValueError("Request an explicit update: provide approved_metadata or set renumber_existing=true.")
        if approved_metadata == {} and not renumber_existing:
            raise ValueError("approved_metadata is empty and no renumbering was requested.")
        if approved_metadata is not None and not all(
            isinstance(key, str) and isinstance(value, str) for key, value in approved_metadata.items()
        ):
            raise ValueError("approved_metadata must map citation numbers to user-verified reference text.")

        document, rules = Document(source), read_rules(rules_path)
        before = audit_reference_document(document, rules)
        inserted = insert_authoritative_entries(document, rules, approved_metadata or {}) if approved_metadata is not None else 0
        after_insert = audit_reference_document(document, rules)
        renumbered = 0
        if renumber_existing:
            if after_insert["missing_reference_entries"] or after_insert["duplicate_reference_numbers"]:
                raise ValueError("Refusing to renumber until missing and duplicate bibliography entries are resolved.")
            renumbered = renumber_reference_labels(document, rules)
        _save_document(document, output)
        return {
            "output_path": str(output),
            "audit_before": before,
            "audit_after": audit_reference_document(document, rules),
            "inserted_authoritative_entries": inserted,
            "renumbered_entries": renumbered,
        }

    return _execute("update_references_copy", action)


@server.tool(annotations=_WRITE_OUTPUT)
def normalize_floating_figures_copy(
    source_docx: str,
    output_docx: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    horizontal: Horizontal = "center",
    vertical: Vertical = "top",
    relative_to: RelativeTo = "paragraph",
    overwrite: bool = False,
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """规范浮动图片锚点并生成 DOCX 副本，不改变图片数据或题注。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        source = _input_path(source_docx, effective_roots, ".docx")
        output = _output_path(
            output_docx,
            workspace_roots=effective_roots,
            inputs=(source,),
            suffix=".docx",
            overwrite=overwrite,
        )
        count = 0

        def transform(name: str, data: bytes) -> bytes:
            nonlocal count
            if not (name.startswith("word/") and name.endswith(".xml")):
                return data
            root = _parse_xml(data)
            anchors = root.xpath("//wp:anchor", namespaces=FIGURE_NS)
            if not anchors:
                return data
            for anchor in anchors:
                configure_anchor(anchor, horizontal, vertical, relative_to)
            count += len(anchors)
            return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)

        _rewrite_docx(source, output, transform)
        return {"output_path": str(output), "normalized_anchor_count": count}

    return _execute("normalize_floating_figures_copy", action)


@server.tool(annotations=_WRITE_OUTPUT)
def mark_word_fields_for_update_copy(
    source_docx: str,
    output_docx: str,
    workspace_roots: Annotated[list[Path], Resolve(_resolve_workspace_roots)],
    overwrite: bool = False,
    workspace_root: str | None = None,
) -> dict[str, Any]:
    """在 DOCX 副本中标记 Word 打开文件时更新目录、交叉引用和页码字段。"""

    def action() -> dict[str, Any]:
        effective_roots, _ = _effective_workspace_roots(workspace_roots, workspace_root)
        source = _input_path(source_docx, effective_roots, ".docx")
        output = _output_path(
            output_docx,
            workspace_roots=effective_roots,
            inputs=(source,),
            suffix=".docx",
            overwrite=overwrite,
        )
        counts = {"toc": 0, "ref": 0, "pageref": 0}
        word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        namespace = {"w": word_namespace}
        with zipfile.ZipFile(source, "r") as archive:
            if "word/settings.xml" not in archive.namelist():
                raise ValueError("The DOCX archive has no word/settings.xml part.")

        def transform(name: str, data: bytes) -> bytes:
            nonlocal counts
            if not (name.startswith("word/") and name.endswith(".xml")):
                return data
            root = _parse_xml(data)
            current = field_counts(root)
            counts = {key: counts[key] + current[key] for key in counts}
            if name != "word/settings.xml":
                return data
            element = root.find("w:updateFields", namespace)
            if element is None:
                element = etree.SubElement(root, f"{{{word_namespace}}}updateFields")
            element.set(f"{{{word_namespace}}}val", "true")
            return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)

        _rewrite_docx(source, output, transform)
        return {"output_path": str(output), "field_counts": counts, "updates_on_open": True}

    return _execute("mark_word_fields_for_update_copy", action)


def main() -> None:
    if sys.version_info[:3] != (3, 13, 15):
        actual = ".".join(map(str, sys.version_info[:3]))
        raise SystemExit(f"This MCP server requires Python 3.13.15; received {actual}. Run it with uv.")
    server.run()


if __name__ == "__main__":
    main()
